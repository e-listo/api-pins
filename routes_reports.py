"""Modul laporan, dua tingkat:
1. Laporan internal biasa -- bisa diakses semua role login, untuk monitoring harian.
2. Laporan resmi/tersertifikasi -- HANYA superadmin/Verifikator/Verifikator Utama,
   dipakai untuk rekonsiliasi SIMBARA (BPKAD Pemkot) atau bahan pemeriksaan/audit.
   Laporan resmi punya blok pengesahan (nama, NIP, jabatan, waktu) tercetak di file.
"""

import io
from datetime import datetime

from flask import Blueprint, request, jsonify, g, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from models import Barang, Transaksi, Pegawai
from auth import login_required, require_role

bp = Blueprint("reports", __name__, url_prefix="/api/v1/reports")


def _kirim_excel(wb: Workbook, nama_file: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=nama_file,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/stok", methods=["GET"])
@login_required
def laporan_stok_internal():
    """Laporan internal biasa -- untuk monitoring harian, TIDAK punya status
    resmi/tersertifikasi. Bisa dicetak siapa saja yang login.
    """
    items = Barang.query.order_by(Barang.kategori.asc(), Barang.nama.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Stok"
    ws.append(["Laporan Stok Barang (Internal)"])
    ws.append([f"Dicetak: {datetime.utcnow().strftime('%d-%m-%Y %H:%M')} UTC"])
    ws.append([])
    header = ["Kode Aset", "Nama", "Kategori", "Satuan", "Stok", "Stok Minimum", "Harga Satuan"]
    ws.append(header)
    for cell in ws[4]:
        cell.font = Font(bold=True)

    for b in items:
        ws.append([b.kode_aset, b.nama, b.kategori, b.satuan, b.stok, b.stok_minimum, b.harga_satuan])

    return _kirim_excel(wb, "laporan_stok_internal.xlsx")


@bp.route("/resmi/stok", methods=["GET"])
@require_role("superadmin", "Verifikator", "Verifikator Utama")
def laporan_stok_resmi():
    """Laporan resmi/tersertifikasi untuk keperluan rekon SIMBARA BPKAD atau
    bahan pemeriksaan/audit. Menyertakan blok pengesahan verifikator.
    """
    verifikator = Pegawai.query.get(g.current_user["sub"])
    items = Barang.query.order_by(Barang.kategori.asc(), Barang.nama.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Resmi Stok"

    ws.append(["PEMERINTAH KOTA YOGYAKARTA"])
    ws.append(["DINAS PEKERJAAN UMUM, PERUMAHAN, DAN KAWASAN PERMUKIMAN"])
    ws.append(["LAPORAN STOK BARANG -- DOKUMEN RESMI/TERSERTIFIKASI"])
    ws.append(["Untuk keperluan rekonsiliasi SIMBARA / bahan pemeriksaan"])
    ws.append([])
    for row in range(1, 4):
        ws.cell(row=row, column=1).font = Font(bold=True)

    header_row = 6
    header = ["Kode Aset", "Nama", "Kategori", "Satuan", "Stok", "Stok Minimum", "Harga Satuan", "Total Nilai"]
    ws.append(header)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    total_keseluruhan = 0
    for b in items:
        nilai = (b.harga_satuan or 0) * (b.stok or 0)
        total_keseluruhan += nilai
        ws.append([b.kode_aset, b.nama, b.kategori, b.satuan, b.stok, b.stok_minimum, b.harga_satuan, nilai])

    ws.append([])
    ws.append(["", "", "", "", "", "", "TOTAL NILAI PERSEDIAAN", total_keseluruhan])

    ws.append([])
    ws.append([])
    ws.append(["Dokumen ini disahkan secara elektronik oleh:"])
    ws.append([f"Nama       : {verifikator.nama}"])
    ws.append([f"NIP        : {verifikator.nip}"])
    ws.append([f"Jabatan    : {verifikator.jabatan or '-'}"])
    ws.append([f"Peran      : {verifikator.role}"])
    ws.append([f"Waktu cetak: {datetime.utcnow().strftime('%d-%m-%Y %H:%M')} UTC"])

    nama_file = f"laporan_resmi_stok_{datetime.utcnow().strftime('%Y%m%d%H%M')}.xlsx"
    return _kirim_excel(wb, nama_file)


@bp.route("/resmi/transaksi", methods=["GET"])
@require_role("superadmin", "Verifikator", "Verifikator Utama")
def laporan_transaksi_resmi():
    """Laporan resmi transaksi (masuk/keluar/mutasi) untuk periode tertentu.
    Query params: ?dari=YYYY-MM-DD&sampai=YYYY-MM-DD
    """
    verifikator = Pegawai.query.get(g.current_user["sub"])
    dari = request.args.get("dari")
    sampai = request.args.get("sampai")

    q = Transaksi.query
    if dari:
        q = q.filter(Transaksi.tanggal >= dari)
    if sampai:
        q = q.filter(Transaksi.tanggal <= sampai)
    items = q.order_by(Transaksi.tanggal.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Resmi Transaksi"

    ws.append(["PEMERINTAH KOTA YOGYAKARTA"])
    ws.append(["DINAS PEKERJAAN UMUM, PERUMAHAN, DAN KAWASAN PERMUKIMAN"])
    periode = f"Periode: {dari or 'awal'} s/d {sampai or 'sekarang'}"
    ws.append(["LAPORAN TRANSAKSI BARANG -- DOKUMEN RESMI/TERSERTIFIKASI"])
    ws.append([periode])
    ws.append([])
    for row in range(1, 4):
        ws.cell(row=row, column=1).font = Font(bold=True)

    header_row = 6
    ws.append(["Nomor Transaksi", "Tanggal", "Tipe", "Jumlah", "Harga Satuan", "Total Nilai", "No. SPK", "Keterangan"])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for t in items:
        ws.append([
            t.nomor_transaksi, str(t.tanggal), t.tipe, t.jumlah,
            float(t.harga_satuan or 0), float(t.total_nilai or 0),
            t.no_spk, t.keterangan,
        ])

    ws.append([])
    ws.append([])
    ws.append(["Dokumen ini disahkan secara elektronik oleh:"])
    ws.append([f"Nama       : {verifikator.nama}"])
    ws.append([f"NIP        : {verifikator.nip}"])
    ws.append([f"Jabatan    : {verifikator.jabatan or '-'}"])
    ws.append([f"Peran      : {verifikator.role}"])
    ws.append([f"Waktu cetak: {datetime.utcnow().strftime('%d-%m-%Y %H:%M')} UTC"])

    nama_file = f"laporan_resmi_transaksi_{datetime.utcnow().strftime('%Y%m%d%H%M')}.xlsx"
    return _kirim_excel(wb, nama_file)
